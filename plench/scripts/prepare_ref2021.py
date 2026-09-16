#!/usr/bin/env python3
"""Download and prepare the official REF 2021 UoA 11 natural LLP bags."""

from __future__ import annotations

import argparse
import collections
import hashlib
import json
import math
import re
import shutil
import sys
import urllib.request
import warnings
from pathlib import Path
from typing import Any, Iterable

import numpy as np

from ..data.ref2021 import CLASS_NAMES, NUM_CLASSES


RESULTS_URL = "https://results2021.ref.ac.uk/profiles/export-all"
OUTPUTS_URL = "https://results2021.ref.ac.uk/outputs/export-all"
PREPROCESSING_VERSION = "ref2021-uoa11-natural-bags-v1"


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _atomic_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(json.dumps(value, indent=2, sort_keys=True), encoding="utf-8")
    temporary.replace(path)


def _download_export(url: str, destination: Path) -> None:
    """Start the official asynchronous export and download the finished xlsx."""
    destination.parent.mkdir(parents=True, exist_ok=True)
    opener = urllib.request.build_opener(urllib.request.HTTPCookieProcessor())
    request = urllib.request.Request(url, data=b"", method="POST")
    with opener.open(request, timeout=120) as response:
        response.read()
    temporary = destination.with_suffix(destination.suffix + ".part")
    with opener.open(url, timeout=600) as response, temporary.open("wb") as output:
        shutil.copyfileobj(response, output)
    if temporary.stat().st_size < 10_000:
        raise RuntimeError(f"official export at {url} is unexpectedly small")
    temporary.replace(destination)


def _download_sources(raw_dir: Path, force: bool) -> tuple[Path, Path]:
    results = raw_dir / "ref2021_results_all.xlsx"
    outputs = raw_dir / "ref2021_outputs_all.xlsx"
    for url, path in ((RESULTS_URL, results), (OUTPUTS_URL, outputs)):
        if path.is_file() and not force:
            print(f"Using existing official export: {path}")
        else:
            print(f"Downloading official export: {url}")
            _download_export(url, path)
    return results, outputs


def _rows_from_xlsx(path: Path, header_first_cell: str) -> Iterable[dict[str, Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ImportError(
            "REF2021 preprocessing requires openpyxl; install "
            "plench/requirements-ref2021.txt"
        ) from exc
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook[workbook.sheetnames[0]]
    # The official profiles export currently declares a stale A1:R8 worksheet
    # dimension even though it contains 7,560 rows. openpyxl read-only mode
    # trusts that declaration unless dimensions are reset explicitly.
    if worksheet.max_row < 100:
        worksheet.reset_dimensions()
    iterator = worksheet.iter_rows(values_only=True)
    header: list[str] | None = None
    for row in iterator:
        values = ["" if value is None else str(value).strip() for value in row]
        if values and values[0] == header_first_cell:
            header = values
            break
    if header is None:
        raise ValueError(f"cannot find {header_first_cell!r} header in {path}")
    for row in iterator:
        values = ["" if value is None else str(value).strip() for value in row]
        if not any(values):
            continue
        values.extend([""] * (len(header) - len(values)))
        yield dict(zip(header, values[: len(header)]))


def _submission_suffix(value: Any) -> str:
    text = str(value or "").strip()
    return text if text else "1"


def _profile_bag_id(row: dict[str, Any], uoa: int) -> str:
    return (
        f"{str(row['Institution code (UKPRN)']).strip()}_{uoa}_"
        f"{_submission_suffix(row.get('Multiple submission letter'))}"
    )


def _output_bag_id(row: dict[str, Any], uoa: int) -> str:
    return (
        f"{str(row['Institution UKPRN code']).strip()}_{uoa}_"
        f"{_submission_suffix(row.get('Multiple submission letter'))}"
    )


def _yes(value: Any) -> bool:
    return str(value or "").strip().lower() in {"yes", "y", "true", "1"}


def _normalise_doi(value: Any) -> str:
    text = str(value or "").strip().lower()
    text = re.sub(r"^(?:https?://(?:dx\.)?doi\.org/|doi:\s*)", "", text)
    return text.rstrip("./ ")


def _normalise_title(value: Any) -> str:
    text = str(value or "").lower()
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return " ".join(text.split())


def _text(row: dict[str, Any]) -> str:
    parts = [f"Title: {row['Title']}", f"Type: {row['Output type']}"]
    venue = row.get("Volume title") or row.get("Publisher") or row.get("Place")
    if venue:
        parts.append(f"Venue: {venue}")
    if row.get("Year"):
        parts.append(f"Year: {row['Year']}")
    return "\n".join(parts)


def _integer_counts(proportions: np.ndarray, total: int) -> np.ndarray:
    raw = proportions * int(total)
    counts = np.floor(raw).astype(np.int64)
    remainder = int(total) - int(counts.sum())
    if remainder:
        order = np.argsort(-(raw - counts), kind="stable")
        counts[order[:remainder]] += 1
    if int(counts.sum()) != int(total):
        raise AssertionError("largest-remainder class counts do not sum to bag size")
    return counts


def _split_score(
    assignment: dict[str, str],
    bags: list[dict[str, Any]],
    duplicate_dois: list[set[str]],
) -> float:
    splits = ("train", "val", "test")
    global_prop = np.stack([bag["class_proportions"] for bag in bags]).mean(axis=0)
    global_log_size = np.mean([math.log1p(bag["n_instances"]) for bag in bags])
    balance = 0.0
    for split in splits:
        selected = [bag for bag in bags if assignment[bag["bag_id"]] == split]
        proportions = np.stack([bag["class_proportions"] for bag in selected])
        balance += float(np.square(proportions.mean(axis=0) - global_prop).sum())
        balance += 0.02 * (
            np.mean([math.log1p(bag["n_instances"]) for bag in selected])
            - global_log_size
        ) ** 2
    leakage = sum(
        len({assignment[bag_id] for bag_id in group}) - 1
        for group in duplicate_dois
    )
    return 0.25 * float(leakage) + balance


def _make_split(
    bags: list[dict[str, Any]],
    doi_to_bags: dict[str, set[str]],
    seed: int,
) -> tuple[dict[str, str], dict[str, Any]]:
    bag_ids = np.asarray(sorted(bag["bag_id"] for bag in bags), dtype=object)
    total = len(bag_ids)
    train_count = int(round(total * 0.60))
    val_count = int(round(total * 0.20))
    test_count = total - train_count - val_count
    duplicate_groups = [value for value in doi_to_bags.values() if len(value) > 1]
    rng = np.random.default_rng(int(seed))
    best_assignment: dict[str, str] | None = None
    best_score = float("inf")
    # Fixed-seed candidate optimization balances proportions/sizes and minimizes
    # unavoidable DOI leakage without ever splitting a natural bag.
    for _ in range(4096):
        order = rng.permutation(bag_ids)
        assignment = {
            str(bag_id): (
                "train"
                if index < train_count
                else "val"
                if index < train_count + val_count
                else "test"
            )
            for index, bag_id in enumerate(order)
        }
        score = _split_score(assignment, bags, duplicate_groups)
        if score < best_score:
            best_score = score
            best_assignment = assignment
    assert best_assignment is not None
    counts = collections.Counter(best_assignment.values())
    assert counts == {"train": train_count, "val": val_count, "test": test_count}
    train = {bag for bag, split in best_assignment.items() if split == "train"}
    val = {bag for bag, split in best_assignment.items() if split == "val"}
    test = {bag for bag, split in best_assignment.items() if split == "test"}
    assert train.isdisjoint(val)
    assert train.isdisjoint(test)
    assert val.isdisjoint(test)
    leaking = {
        doi: {
            "bags": sorted(group),
            "splits": sorted({best_assignment[bag] for bag in group}),
        }
        for doi, group in doi_to_bags.items()
        if len(group) > 1 and len({best_assignment[bag] for bag in group}) > 1
    }
    diagnostics = {
        "strategy": "fixed_seed_proportion_size_balanced_with_doi_leakage_penalty",
        "candidate_assignments": 4096,
        "score": best_score,
        "target_bag_counts": {
            "train": train_count,
            "val": val_count,
            "test": test_count,
        },
        "duplicate_doi_groups": len(duplicate_groups),
        "cross_split_duplicate_doi_groups": len(leaking),
        "cross_split_duplicate_dois": leaking,
    }
    return best_assignment, diagnostics


def _fixed_list_array(pa: Any, matrix: np.ndarray, value_type: Any) -> Any:
    values = pa.array(matrix.reshape(-1), type=value_type)
    return pa.FixedSizeListArray.from_arrays(values, matrix.shape[1])


def prepare(data_root: Path, uoa: int, seed: int, force_download: bool) -> dict[str, Any]:
    if int(uoa) != 11:
        raise ValueError("this benchmark adapter is intentionally fixed to REF2021 UoA 11")
    try:
        import pyarrow as pa
        import pyarrow.parquet as pq
    except ImportError as exc:
        raise ImportError(
            "REF2021 preprocessing requires pyarrow; install "
            "plench/requirements-ref2021.txt"
        ) from exc

    raw_dir = data_root / "raw"
    processed = data_root / "processed"
    splits_dir = data_root / "splits"
    results_path, outputs_path = _download_sources(raw_dir, force_download)

    profiles = [
        row
        for row in _rows_from_xlsx(results_path, "Institution code (UKPRN)")
        if str(row.get("Unit of assessment number")) == str(uoa)
        and str(row.get("Profile")).strip() == "Outputs"
    ]
    outputs = [
        row
        for row in _rows_from_xlsx(outputs_path, "Institution UKPRN code")
        if str(row.get("Unit of assessment number")) == str(uoa)
    ]
    if not profiles or not outputs:
        raise ValueError("official exports contain no UoA 11 Outputs Profile data")

    profile_by_bag: dict[str, dict[str, Any]] = {}
    raw_sum_deviations: list[float] = []
    for row in profiles:
        bag_id = _profile_bag_id(row, uoa)
        if bag_id in profile_by_bag:
            raise ValueError(f"duplicate official Outputs Profile for {bag_id}")
        vector = np.asarray(
            [
                float(row.get("Unclassified") or 0),
                float(row.get("1*") or 0),
                float(row.get("2*") or 0),
                float(row.get("3*") or 0),
                float(row.get("4*") or 0),
            ],
            dtype=np.float64,
        ) / 100.0
        raw_sum = float(vector.sum())
        deviation = abs(raw_sum - 1.0)
        raw_sum_deviations.append(deviation)
        if deviation > 0.02:
            warnings.warn(
                f"{bag_id} official Outputs Profile sums to {raw_sum:.6f}; inspect source",
                stacklevel=2,
            )
        if raw_sum <= 0 or (vector < 0).any():
            raise ValueError(f"{bag_id}: invalid official Outputs Profile")
        row = dict(row)
        row["normalised_proportions"] = (vector / raw_sum).astype(np.float32)
        row["raw_proportion_sum"] = raw_sum
        profile_by_bag[bag_id] = row

    output_groups: dict[str, list[dict[str, Any]]] = collections.defaultdict(list)
    for row in outputs:
        output_groups[_output_bag_id(row, uoa)].append(row)
    unmatched_profiles = sorted(set(profile_by_bag) - set(output_groups))
    unmatched_outputs = sorted(set(output_groups) - set(profile_by_bag))
    if unmatched_profiles or unmatched_outputs:
        print("Unmatched official profiles:", unmatched_profiles, file=sys.stderr)
        print("Unmatched output submissions:", unmatched_outputs, file=sys.stderr)
        raise ValueError("REF2021 profile/output matching is not one-to-one")

    instances: list[dict[str, Any]] = []
    bag_rows: list[dict[str, Any]] = []
    doi_to_bags: dict[str, set[str]] = collections.defaultdict(set)
    title_to_bags: dict[str, set[str]] = collections.defaultdict(set)
    output_types: collections.Counter[str] = collections.Counter()
    double_rows = 0
    reserve_rows = 0
    for bag_id in sorted(profile_by_bag):
        profile = profile_by_bag[bag_id]
        rows = sorted(
            output_groups[bag_id],
            key=lambda value: (str(value.get("REF2ID")), str(value.get("Title"))),
        )
        ref2id_occurrences: collections.Counter[str] = collections.Counter()
        for row in rows:
            proposed_double = _yes(row.get("Propose double weighting"))
            reserve = _yes(row.get("Is reserve output"))
            double_rows += int(proposed_double)
            reserve_rows += int(reserve)
            # The public output table records proposed, not accepted, double
            # weighting. UoA 11 has zero affected records, so effective weight is
            # unambiguously one for every instance.
            instance_weight = 1.0
            ref2id = str(row.get("REF2ID") or "").strip()
            if not ref2id:
                ref2id = hashlib.sha256(
                    f"{bag_id}\0{row.get('Title')}\0{row.get('Year')}".encode()
                ).hexdigest()[:24]
            ref2id_occurrences[ref2id] += 1
            doi = _normalise_doi(row.get("DOI"))
            normalized_title = _normalise_title(row.get("Title"))
            if doi:
                doi_to_bags[doi].add(bag_id)
            if normalized_title:
                title_to_bags[normalized_title].add(bag_id)
            output_types[str(row.get("Output type") or "unknown")] += 1
            instances.append(
                {
                    # REF2ID can recur when the same output appears in several
                    # submissions. The composite ID identifies the submitted
                    # output instance without pretending those rows are unique
                    # scholarly works.
                    "instance_id": (
                        f"{bag_id}:{ref2id}:{ref2id_occurrences[ref2id]:02d}"
                    ),
                    "bag_id": bag_id,
                    "institution_code": str(row["Institution UKPRN code"]),
                    "institution_name": str(row["Institution name"]),
                    "uoa_code": int(uoa),
                    "multiple_submission_id": _submission_suffix(
                        row.get("Multiple submission letter")
                    ),
                    "output_id": ref2id,
                    "title": str(row.get("Title") or ""),
                    "output_type": str(row.get("Output type") or ""),
                    "year": int(float(row["Year"])) if row.get("Year") else None,
                    "doi": doi,
                    "venue": str(
                        row.get("Volume title")
                        or row.get("Publisher")
                        or row.get("Place")
                        or ""
                    ),
                    "text": _text(row),
                    "instance_weight": instance_weight,
                    "instance_label": -1,
                    "propose_double_weighting": proposed_double,
                    "is_reserve_output": reserve,
                    "feature_index": len(instances),
                }
            )
        vector = np.asarray(profile["normalised_proportions"], dtype=np.float32)
        counts = _integer_counts(vector.astype(np.float64), len(rows))
        bag_rows.append(
            {
                "bag_id": bag_id,
                "institution_code": str(profile["Institution code (UKPRN)"]),
                "institution_name": str(profile["Institution name"]),
                "uoa_code": int(uoa),
                "multiple_submission_id": _submission_suffix(
                    profile.get("Multiple submission letter")
                ),
                "n_instances": len(rows),
                "effective_n_instances": float(len(rows)),
                "class_proportions": vector,
                "inferred_class_counts": counts,
                "raw_proportion_sum": float(profile["raw_proportion_sum"]),
            }
        )

    if double_rows or reserve_rows:
        raise ValueError(
            "UoA 11 unexpectedly contains double-weight/reserve records; accepted "
            "multiplicity must be resolved before preprocessing"
        )

    assignment, split_diagnostics = _make_split(bag_rows, doi_to_bags, seed)
    for row in bag_rows:
        row["split"] = assignment[row["bag_id"]]
    split_manifest = {
        "dataset": "ref2021_uoa11",
        "seed": int(seed),
        "ratios": {"train": 0.6, "val": 0.2, "test": 0.2},
        "bags": {
            split: sorted(bag for bag, value in assignment.items() if value == split)
            for split in ("train", "val", "test")
        },
        "diagnostics": split_diagnostics,
    }
    split_path = splits_dir / f"ref2021_uoa11_seed{seed}.json"
    _atomic_json(split_path, split_manifest)

    processed.mkdir(parents=True, exist_ok=True)
    instance_table = pa.Table.from_pylist(instances)
    pq.write_table(instance_table, processed / "instances.parquet", compression="zstd")
    proportions_matrix = np.stack([row.pop("class_proportions") for row in bag_rows])
    counts_matrix = np.stack([row.pop("inferred_class_counts") for row in bag_rows])
    bag_table = pa.table(
        {
            **{
                key: pa.array([row[key] for row in bag_rows])
                for key in bag_rows[0]
            },
            "class_proportions": _fixed_list_array(
                pa, proportions_matrix.astype(np.float32), pa.float32()
            ),
            **{
                f"p_{name}": pa.array(
                    proportions_matrix[:, index].astype(np.float32),
                    type=pa.float32(),
                )
                for index, name in enumerate(CLASS_NAMES)
            },
            "inferred_class_counts": _fixed_list_array(
                pa, counts_matrix.astype(np.int64), pa.int64()
            ),
        }
    )
    pq.write_table(bag_table, processed / "bags.parquet", compression="zstd")

    bag_sizes = np.asarray([row["n_instances"] for row in bag_rows], dtype=np.int64)
    all_proportions = proportions_matrix.astype(np.float64)
    doi_values = collections.Counter(instance["doi"] for instance in instances if instance["doi"])
    title_values = collections.Counter(
        _normalise_title(instance["title"]) for instance in instances if instance["title"]
    )
    split_counts = collections.Counter(row["split"] for row in bag_rows)
    stats = {
        "total_bags": len(bag_rows),
        "total_official_output_records": len(instances),
        "total_unique_outputs": len({instance["output_id"] for instance in instances}),
        "duplicate_output_id_extra_rows": (
            len(instances) - len({instance["output_id"] for instance in instances})
        ),
        "effective_total_outputs": float(sum(row["effective_n_instances"] for row in bag_rows)),
        "classes": NUM_CLASSES,
        "class_names": CLASS_NAMES,
        "bag_size": {
            "min": int(bag_sizes.min()),
            "q25": float(np.quantile(bag_sizes, 0.25)),
            "median": float(np.median(bag_sizes)),
            "mean": float(bag_sizes.mean()),
            "q75": float(np.quantile(bag_sizes, 0.75)),
            "max": int(bag_sizes.max()),
            "histogram": {
                str(size): int(count)
                for size, count in sorted(collections.Counter(bag_sizes.tolist()).items())
            },
        },
        "split_bags": {key: int(split_counts[key]) for key in ("train", "val", "test")},
        "class_proportions": {
            name: {
                "mean": float(all_proportions[:, index].mean()),
                "std": float(all_proportions[:, index].std()),
                "min": float(all_proportions[:, index].min()),
                "max": float(all_proportions[:, index].max()),
            }
            for index, name in enumerate(CLASS_NAMES)
        },
        "mean_official_class_proportions": all_proportions.mean(axis=0).tolist(),
        "std_official_class_proportions": all_proportions.std(axis=0).tolist(),
        "outputs_with_doi": int(sum(bool(instance["doi"]) for instance in instances)),
        "outputs_without_doi": int(sum(not instance["doi"] for instance in instances)),
        "doi_coverage": float(sum(bool(instance["doi"]) for instance in instances) / len(instances)),
        "duplicate_doi_values": int(sum(count > 1 for count in doi_values.values())),
        "duplicate_doi_extra_rows": int(sum(count - 1 for count in doi_values.values() if count > 1)),
        "cross_bag_duplicate_doi_values": int(sum(len(value) > 1 for value in doi_to_bags.values())),
        "duplicate_normalized_title_values": int(sum(count > 1 for count in title_values.values())),
        "cross_bag_duplicate_normalized_titles": int(sum(len(value) > 1 for value in title_to_bags.values())),
        "output_type_distribution": dict(sorted(output_types.items())),
        "double_weighted_related_records": int(double_rows),
        "reserve_output_related_records": int(reserve_rows),
        "unmatched_outputs": unmatched_outputs,
        "unmatched_profiles": unmatched_profiles,
        "official_profile_count": len(profiles),
        "constructed_bag_count": len(bag_rows),
        "raw_proportion_sum_deviation": {
            "max": float(max(raw_sum_deviations)),
            "mean": float(np.mean(raw_sum_deviations)),
        },
        "split_diagnostics": split_diagnostics,
        "first_five_bags": [
            {
                "institution_name": row["institution_name"],
                "bag_id": row["bag_id"],
                "n_instances": row["n_instances"],
                "proportion": proportions_matrix[index].tolist(),
            }
            for index, row in enumerate(bag_rows[:5])
        ],
    }
    _atomic_json(processed / "ref2021_uoa11_stats.json", stats)
    split_manifest_sha256 = _sha256(split_path)
    source_hashes = {
        "results": _sha256(results_path),
        "outputs": _sha256(outputs_path),
    }
    preprocessing_hash = hashlib.sha256(
        json.dumps(
            {
                "version": PREPROCESSING_VERSION,
                "sources": source_hashes,
                "split_manifest_sha256": split_manifest_sha256,
                "class_names": CLASS_NAMES,
                "uoa": 11,
            },
            sort_keys=True,
        ).encode()
    ).hexdigest()
    metadata = {
        "dataset": "ref2021_uoa11",
        "uoa": 11,
        "uoa_name": "Computer Science and Informatics",
        "class_names": CLASS_NAMES,
        "num_classes": NUM_CLASSES,
        "has_instance_labels": False,
        "natural_bags": True,
        "variable_bag_size": True,
        "profile": "Outputs",
        "preprocessing_version": PREPROCESSING_VERSION,
        "preprocessing_hash": preprocessing_hash,
        "split_seed": int(seed),
        "split_manifest": str(split_path.relative_to(data_root)),
        "source_files": {
            "results": {
                "path": str(results_path.relative_to(data_root)),
                "url": RESULTS_URL,
                "sha256": source_hashes["results"],
            },
            "outputs": {
                "path": str(outputs_path.relative_to(data_root)),
                "url": OUTPUTS_URL,
                "sha256": source_hashes["outputs"],
            },
        },
        "bag_id_definition": "{institution_ukprn}_11_{multiple_submission_letter_or_1}",
        "split_manifest_sha256": split_manifest_sha256,
        "instance_weighting": {
            "supported": True,
            "uoa11_all_weights": 1.0,
            "reason": "official UoA11 export has no proposed-double-weight or reserve rows",
        },
        "inferred_class_counts": {
            "used_by_training": False,
            "purpose": "diagnostics and count-based methods only",
            "method": "largest remainder from rounded normalized Outputs Profile and n_instances",
        },
    }
    _atomic_json(processed / "metadata.json", metadata)
    print(json.dumps(stats, indent=2, sort_keys=True))
    return stats


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--uoa", type=int, default=11)
    parser.add_argument("--data-root", type=Path, default=Path("plench/data/ref2021_uoa11"))
    parser.add_argument("--seed", type=int, default=42)
    parser.add_argument("--force-download", action="store_true")
    args = parser.parse_args()
    prepare(args.data_root.expanduser().resolve(), args.uoa, args.seed, args.force_download)


if __name__ == "__main__":
    main()
